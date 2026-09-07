from __future__ import annotations

import json

import httpx
import pytest
from docx import Document

from manual_generator.copyright_documents import build_code
from manual_generator.copyright_schemas import Business, Registration, Section
from manual_generator.copyright_service import invalidate, require_confirmations
from manual_generator.copyright_sources import (
    digest,
    scan_sources,
    source_pages,
    verify_sources,
    wrap_code,
)
from manual_generator.llm import ModelOptions, OpenAICompatibleClient
from manual_generator.models import CopyrightCase


def test_scan_excludes_sensitive_dependencies_and_preserves_line_numbers(tmp_path):
    (tmp_path / "main.py").write_text("# demo\n\ndef calculate(a, b):\n    return a + b\n")
    (tmp_path / "secret.py").write_text('API_KEY = "sensitive-value"')
    (tmp_path / ".env").write_text("SECRET=private")
    (tmp_path / "node_modules").mkdir()
    (tmp_path / "node_modules" / "dependency.js").write_text("external()")
    inventory = scan_sources(tmp_path)
    assert [file["path"] for file in inventory["files"]] == ["main.py"]
    assert "sensitive-value" not in json.dumps(inventory)
    assert len(inventory["excluded"]) == 2
    pages = source_pages(tmp_path, inventory, ["main.py"])
    rows = pages["volumes"][0]["pages"][0]["rows"]
    assert [row["line"] for row in rows] == [1, 3, 4]
    assert "3: def calculate" in inventory["files"][0]["excerpt"]


@pytest.mark.parametrize("count,volume_lengths", [(49, [1]), (50, [1]), (51, [2]), (2999, [60]), (3000, [60]), (3001, [30, 30]), (5000, [30, 30])])
def test_real_code_pagination(tmp_path, count, volume_lengths):
    (tmp_path / "main.py").write_text("\n".join(f"line_{i} = {i}" for i in range(count)))
    inventory = scan_sources(tmp_path)
    result = source_pages(tmp_path, inventory, ["main.py"])
    assert [len(volume["pages"]) for volume in result["volumes"]] == volume_lengths
    final_row = result["volumes"][-1]["pages"][-1]["rows"][-1]
    assert final_row["line"] == count
    for volume in result["volumes"]:
        for page in volume["pages"]:
            assert len(page["rows"]) == 50 or page["number"] == result["total_pages"]


def test_code_unicode_and_original_content():
    original = "\t" + "中文" * 80 + "a" * 100
    wrapped = wrap_code(original)
    assert "".join(wrapped) == original.expandtabs(4)
    assert len(wrapped) > 2
    assert all(len(line) <= 88 for line in wrapped)


@pytest.mark.parametrize("secret", ['DB_URL = "postgres://alice:privatepassword@database/app"', 'token = "ghp_012345678901234567890123"', 'headers = {"Authorization": "Bearer abcdefghijklmnopqr"}'])
def test_secret_url_and_token_exclusion(tmp_path, secret):
    (tmp_path / "main.py").write_text("print(1)")
    (tmp_path / "config.py").write_text(secret)
    inventory = scan_sources(tmp_path)
    assert [file["path"] for file in inventory["files"]] == ["main.py"]
    assert "privatepassword" not in json.dumps(inventory)


def test_code_rejects_path_escape_duplicate_and_changed_sources(tmp_path):
    source = tmp_path / "main.py"
    source.write_text("print(1)")
    inventory = scan_sources(tmp_path)
    for paths in (["../main.py"], ["main.py", "main.py"], []):
        with pytest.raises(ValueError):
            source_pages(tmp_path, inventory, paths)
    source.write_text("print(2)")
    with pytest.raises(ValueError, match="变化"):
        source_pages(tmp_path, inventory, ["main.py"])
    with pytest.raises(ValueError, match="变化"):
        verify_sources(tmp_path, inventory)


def test_stage_invalidates_only_dependents():
    case = CopyrightCase(revision=8, status="idle", data={"business": {"x": 1}, "registration": {"x": 2}, "sources": {"x": 3}, "drafts": {}, "draft_checkpoint": {}, "reviewed_digest": "old"}, confirmations={key: {"digest": digest({"x": i})} for i, key in enumerate(["business", "registration", "sources"], 1)})
    require_confirmations(case, "sources")
    invalidate(case, "registration")
    assert case.revision == 9
    assert list(case.confirmations) == ["business"]
    assert "drafts" not in case.data
    assert "draft_checkpoint" not in case.data
    with pytest.raises(ValueError):
        require_confirmations(case, "sources")
    case.data = {**case.data, "business": {"x": 9}}
    with pytest.raises(ValueError):
        require_confirmations(case, "business")


def registration() -> dict:
    values = {key: "测试内容" for key in Registration.model_fields}
    values.update(software_name="测试任务管理软件", short_name="", version="V1.0", owner_type="法人", development_method="单独开发", originality="原创", completion_date="2026-01-01", publication_status="未发表", publication_date="", publication_place="", rights_acquisition="原始取得", ownership_notes="")
    return values


def test_registration_conditional_dates():
    profile = Registration.model_validate(registration())
    assert profile.blockers() == []
    profile.publication_status = "已发表"
    assert profile.blockers()
    profile.publication_date = "2025-12-01"
    assert any("不早于" in error for error in profile.blockers())
    profile.publication_date = "2026-02-01"
    profile.publication_place = "中国"
    assert profile.blockers() == []


def test_docx_code_has_50_rows_and_no_generator_brand(tmp_path):
    path = tmp_path / "code.docx"
    pages = [{"number": 70, "rows": [{"text": f"value_{i} = {i}"} for i in range(50)]}, {"number": 71, "rows": [{"text": "last = True"}]}]
    build_code(path, registration(), pages)
    doc = Document(path)
    assert len(doc.paragraphs[0].text.splitlines()) == 50
    assert doc.paragraphs[1].paragraph_format.page_break_before
    assert "操作手册生成器" not in doc.sections[0].footer.paragraphs[0].text
    assert "测试任务管理软件" in doc.sections[0].header.paragraphs[0].text


def test_package_without_pdf_preserves_docx_and_records_skipped_checks(tmp_path, monkeypatch):
    import zipfile

    from manual_generator.config import Settings
    from manual_generator.copyright_documents import build_package

    monkeypatch.setattr("manual_generator.copyright_documents.get_settings", lambda: Settings(pdf_enabled=False))
    def reject_conversion(*args):
        pytest.fail("PDF conversion must not run while disabled")
    monkeypatch.setattr("manual_generator.copyright_documents._convert_pdf", reject_conversion)
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (workspace / "main.py").write_text("print('hello')\n")
    inventory = scan_sources(workspace)
    chapter = {"title": "软件说明", "paragraphs": ["展示信息。"], "evidence_ids": ["main.py"], "screenshot_ids": []}
    data = {"registration": registration(), "inventory": inventory, "sources": {"paths": ["main.py"]}, "drafts": {"manual": [chapter], "design": [chapter]}, "rules": {}}
    manifest = build_package(None, data, tmp_path / "copyright" / "batches" / "test", {})
    assert [item["kind"] for item in manifest].count("docx") == 3
    assert not any(item["kind"] == "pdf" for item in manifest)
    report = next(item for item in manifest if item["material"] == "validation")
    from pathlib import Path
    checks = json.loads(Path(report["path"]).read_text())["checks"]
    assert all(item["pdf_verification"] == "skipped" and "pages" not in item for item in checks)
    archive = next(item for item in manifest if item["material"] == "package")
    with zipfile.ZipFile(archive["path"]) as bundle:
        assert len(bundle.namelist()) == 4
        assert not any(name.endswith(".pdf") for name in bundle.namelist())


@pytest.mark.asyncio
@pytest.mark.parametrize("protocol", ["chat_completions", "responses"])
async def test_structured_model_repairs_json_and_supports_both_protocols(monkeypatch, protocol):
    calls = []
    valid = {"title": "模块", "paragraphs": ["真实内容"], "evidence_ids": ["main.py"]}
    def handler(request):
        calls.append(json.loads(request.content))
        content = "invalid json" if len(calls) == 1 else json.dumps(valid)
        response = {"choices": [{"message": {"content": content}}]} if protocol == "chat_completions" else {"output_text": content}
        return httpx.Response(200, json=response)
    original = httpx.AsyncClient
    monkeypatch.setattr("manual_generator.llm.httpx.AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs))
    client = OpenAICompatibleClient(ModelOptions("fake-test-key", "https://model.invalid/v1", "test", protocol))
    result = await client.structured_document("生成文档", {"evidence": []}, Section)
    assert result.title == "模块"
    assert len(calls) == 2


@pytest.mark.asyncio
async def test_model_timeout_is_bounded_and_redacted(monkeypatch):
    calls = []
    def handler(request):
        calls.append(request)
        raise httpx.ReadTimeout("secret-value", request=request)
    original = httpx.AsyncClient
    monkeypatch.setattr("manual_generator.llm.httpx.AsyncClient", lambda **kwargs: original(transport=httpx.MockTransport(handler), **kwargs))
    client = OpenAICompatibleClient(ModelOptions("fake", "https://model.invalid/v1", "test", "chat_completions"))
    with pytest.raises(ValueError, match="超时") as error:
        await client.structured_document("test", {}, Business)
    assert len(calls) == 3
    assert "secret-value" not in str(error.value)
