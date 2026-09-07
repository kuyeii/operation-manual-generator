from __future__ import annotations

import json
import re
import unicodedata
import zipfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.shared import Pt, RGBColor
from pypdf import PdfReader

from .config import get_settings
from .copyright_schemas import FIELD_LABELS
from .copyright_sources import digest, source_pages
from .reports import (
    _add_image,
    _add_page_field,
    _configure,
    _convert_pdf,
    _render_and_verify_pdf,
    _set_run_font,
    safe_name,
)


def configure(document: Document, registration: dict) -> None:
    _configure(document)
    for style in document.styles:
        if style.type == 1 or style.type == 2:
            style.font.color.rgb = RGBColor(0, 0, 0)
    header = document.sections[0].header.paragraphs[0]
    header.text = f"{registration['software_name']}  {registration['version']}"
    for run in header.runs:
        run.font.size = Pt(8)
    footer = document.sections[0].footer.paragraphs[0]
    footer.clear()
    _add_page_field(footer)


def application_text(registration: dict, source_lines: int) -> str:
    text = "\n".join(f"{label}：{registration.get(key, '')}" for key, label in FIELD_LABELS.items())
    return text + f"\n源程序量：{source_lines}\n"


def build_narrative(path: Path, title: str, registration: dict, sections: list[dict], screenshots: dict) -> None:
    document = Document()
    configure(document, registration)
    document.add_heading(registration["software_name"], 0)
    document.add_paragraph(registration["version"])
    document.add_heading(title, 1)
    document.add_page_break()
    document.add_heading("目录", 1)
    for index, section in enumerate(sections, 1):
        document.add_paragraph(f"{index}　{section['title']}")
    document.add_page_break()
    for index, section in enumerate(sections, 1):
        document.add_heading(f"{index}　{section['title']}", 1)
        for paragraph in section["paragraphs"]:
            document.add_paragraph(paragraph)
        for number, screenshot_id in enumerate(section.get("screenshot_ids", []), 1):
            screenshot = screenshots[screenshot_id]
            _add_image(document, Path(screenshot["path"]))
            document.add_paragraph(f"图 {index}-{number}　{screenshot['instruction']}")
    document.save(path)


def build_code(path: Path, registration: dict, pages: list[dict]) -> None:
    document = Document()
    configure(document, registration)
    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    normal.paragraph_format.line_spacing = Pt(12)
    normal.paragraph_format.widow_control = False
    normal.font.size = Pt(8)
    for index, page in enumerate(pages):
        paragraph = document.add_paragraph()
        if index:
            paragraph.paragraph_format.page_break_before = True
        for row_index, row in enumerate(page["rows"]):
            run = paragraph.add_run(row["text"])
            _set_run_font(run, "Noto Sans CJK SC")
            run.font.size = Pt(8)
            if row_index < len(page["rows"]) - 1:
                run.add_break(WD_BREAK.LINE)
    # The back volume starts at its source page, not at one.
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    page_number = OxmlElement("w:pgNumType")
    page_number.set(qn("w:start"), str(pages[0]["number"]))
    document.sections[0]._sectPr.append(page_number)
    document.save(path)


def verify_pdf(pdf: Path, preview_dir: Path, expected_pages: int | None = None) -> int:
    pages = _render_and_verify_pdf(pdf, preview_dir)
    if expected_pages is not None and len(pages) != expected_pages:
        raise ValueError(f"{pdf.name}实际为{len(pages)}页，预期{expected_pages}页，未发布")
    text = "\n".join(page.extract_text() for page in PdfReader(pdf).pages)
    if "\ufffd" in text or not text.strip():
        raise ValueError(f"{pdf.name}存在缺字或文本为空")
    return len(pages)


def verify_code_text(pdf: Path, pages: list[dict]) -> None:
    def normalize(text: str) -> str:
        return re.sub(r"\s+", "", unicodedata.normalize("NFKC", text))
    rendered = PdfReader(pdf).pages
    if len(rendered) != len(pages):
        raise ValueError("源码PDF分页与原始分页模型不一致")
    for page, expected in zip(rendered, pages, strict=True):
        text, cursor = normalize(page.extract_text()), 0
        for row in expected["rows"]:
            fragment = normalize(row["text"])
            position = text.find(fragment, cursor)
            if position < 0:
                raise ValueError(f"源码第{expected['number']}页有文字截断或缺字：{row['path']} 第{row['line']}行")
            cursor = position + len(fragment)


def build_package(task, data: dict, batch_dir: Path, screenshots: dict) -> list[dict]:
    """Write an unpublished batch; the caller atomically registers it after all checks pass."""
    output = batch_dir / "正式资料"
    output.mkdir(parents=True, exist_ok=True)
    registration = data["registration"]
    name = safe_name(registration["software_name"])[:60]
    workspace = batch_dir.parents[2] / "workspace"
    codes = source_pages(workspace, data["inventory"], data["sources"]["paths"])
    manifest, checks = [], []

    def add(path: Path, material: str, formal: bool = True):
        manifest.append({"path": str(path), "name": path.name, "kind": path.suffix.lstrip("."), "material": material, "formal": formal, "size": path.stat().st_size})

    application = output / "申请填报信息.txt"
    application.write_text(application_text(registration, data["inventory"]["source_lines"]), encoding="utf-8")
    add(application, "application")
    documents = []
    for key, title in (("manual", "操作手册"), ("design", "技术设计说明书")):
        path = output / f"{name}_{title}.docx"
        build_narrative(path, title, registration, data["drafts"][key], screenshots)
        documents.append((path, key, None, None))
    for volume in codes["volumes"]:
        path = output / f"{name}_源程序({volume['name']}).docx"
        build_code(path, registration, volume["pages"])
        documents.append((path, "source", len(volume["pages"]), volume["pages"]))
    for path, material, expected_pages, expected_rows in documents:
        if not get_settings().pdf_enabled:
            checks.append({"document": path.name, "expected_pages": expected_pages, "pdf_verification": "skipped", "reason": "PDF生成与版面验收已关闭"})
            add(path, material)
            continue
        converted = _convert_pdf(path, batch_dir / "conversion" / path.stem)
        pdf = output / converted.name
        converted.replace(pdf)
        page_count = verify_pdf(pdf, batch_dir / "previews" / path.stem, expected_pages)
        text = re.sub(r"\s+", "", "\n".join(page.extract_text() for page in PdfReader(pdf).pages))
        for value in (registration["software_name"], registration["version"]):
            if re.sub(r"\s+", "", value) not in text:
                raise ValueError(f"{pdf.name}中软件名称或版本缺失")
        if expected_rows is not None:
            verify_code_text(pdf, expected_rows)
        checks.append({"document": path.name, "pages": page_count, "expected_pages": expected_pages, "passed": True})
        add(path, material)
        add(pdf, material)
    evidence = batch_dir / "证据清单.json"
    evidence.write_text(json.dumps({"inventory": data["inventory"], "selection": data["sources"], "source_pages": codes, "chapters": data["drafts"], "screenshots": screenshots}, ensure_ascii=False, indent=2), encoding="utf-8")
    add(evidence, "evidence", False)
    report = batch_dir / "校验报告.json"
    report.write_text(json.dumps({"passed": True, "rules": data["rules"], "checks": checks, "model_review": data.get("review_findings", {}), "manual_review": data.get("review_resolutions", {}), "revision_fingerprint": digest(data), "source_lines": data["inventory"]["source_lines"], "selected_lines": codes["selected_lines"]}, ensure_ascii=False, indent=2), encoding="utf-8")
    add(report, "validation", False)
    archive = batch_dir / f"{name}_软著正式材料.zip"
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as bundle:
        for item in manifest:
            if item["formal"]:
                bundle.write(item["path"], item["name"])
    add(archive, "package")
    return manifest
