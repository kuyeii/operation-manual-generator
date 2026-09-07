from __future__ import annotations

import os
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from .config import get_settings
from .models import Task

DOCUMENT_FONT = "Noto Sans CJK SC"


def build_report(task: Task, task_dir: Path) -> Path:
    report_dir = task_dir / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)
    docx_path = report_dir / f"{safe_name(task.name)}-用户操作手册.docx"
    if get_settings().pdf_enabled:
        docx_path.with_suffix(".pdf").unlink(missing_ok=True)
    (report_dir / "fontconfig.xml").unlink(missing_ok=True)
    shutil.rmtree(report_dir / ".font-cache", ignore_errors=True)
    document = Document()
    _configure(document)
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(120)
    run = title.add_run(task.name)
    _set_run_font(run, DOCUMENT_FONT)
    run.font.size = Pt(28)
    run.bold = True
    subtitle = document.add_paragraph("用户操作手册")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.style = document.styles["Subtitle"]
    document.add_paragraph(f"生成地址：{task.start_url or '-'}").alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_page_break()
    document.add_heading("目录", level=1)
    _add_toc(document, task)
    document.add_page_break()
    document.add_heading("使用环境", level=1)
    document.add_paragraph("本手册由自动化浏览器在 1440×900 桌面视口下生成。界面内容以生成时的测试系统为准。")
    for index, feature in enumerate(task.features, 1):
        if not feature.selected or feature.status != "completed":
            continue
        document.add_heading(f"{index}. {feature.title}", level=1)
        if feature.goal:
            document.add_paragraph(feature.goal)
        if not feature.steps:
            document.add_paragraph(f"状态：{feature.status}。{feature.error or '未记录操作步骤。'}")
            continue
        for step_index, step in enumerate(feature.steps, 1):
            if not step.screenshot or not step.screenshot.included:
                continue
            document.add_heading(f"步骤 {step_index}", level=2)
            document.add_paragraph(step.instruction)
            if step.screenshot and step.screenshot.included and Path(step.screenshot.path).exists():
                _add_image(document, Path(step.screenshot.path))
                caption = document.add_paragraph(f"图 {index}-{step_index}  {step.instruction}")
                caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption.style = document.styles["Caption"]
    failed = [feature for feature in task.features if feature.status not in {"completed", "pending"}]
    if failed:
        document.add_heading("未覆盖与失败项", level=1)
        for feature in failed:
            document.add_paragraph(f"{feature.title}：{feature.error or feature.status}", style="List Bullet")
    document.save(docx_path)
    if get_settings().pdf_enabled:
        _render_and_verify_docx(docx_path, report_dir / "rendered")
    return docx_path


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.page_height, section.page_width = Cm(29.7), Cm(21)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.25)
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = DOCUMENT_FONT
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), DOCUMENT_FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.35
    for name, size, color in (("Title", 28, "1D252C"), ("Heading 1", 18, "1D252C"), ("Heading 2", 13, "276B5D")):
        style = styles[name]
        style.font.name = DOCUMENT_FONT
        style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCUMENT_FONT)
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
    if "Screenshot Caption" not in styles:
        styles.add_style("Screenshot Caption", WD_STYLE_TYPE.PARAGRAPH)
    if "Manual TOC" not in styles:
        toc_style = styles.add_style("Manual TOC", WD_STYLE_TYPE.PARAGRAPH)
        toc_style.font.name = DOCUMENT_FONT
        toc_style.element.rPr.rFonts.set(qn("w:eastAsia"), DOCUMENT_FONT)
        toc_style.font.size = Pt(11)
        toc_style.font.color.rgb = RGBColor.from_string("39423E")
        toc_style.paragraph_format.space_after = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("操作手册生成器 · ")
    _add_page_field(footer)


def _add_toc(document: Document, task: Task) -> None:
    document.add_paragraph("使用环境", style="Manual TOC")
    for index, feature in enumerate((item for item in task.features if item.selected), 1):
        document.add_paragraph(f"{index}. {feature.title}", style="Manual TOC")
    if any(feature.status not in {"completed", "pending"} for feature in task.features):
        document.add_paragraph("未覆盖与失败项", style="Manual TOC")


def _set_run_font(run, name: str) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def _add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def _add_image(document: Document, path: Path) -> None:
    with Image.open(path) as image:
        ratio = image.height / image.width
    width_cm = min(16.5, 21 / ratio)
    document.add_picture(str(path), width=Cm(width_cm), height=Cm(width_cm * ratio))
    document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def _project_font_dir() -> Path:
    configured = os.environ.get("MANUAL_GENERATOR_FONT_DIR")
    return Path(configured) if configured else Path(__file__).resolve().parents[2] / ".fonts"


def _render_and_verify_docx(docx_path: Path, output_dir: Path) -> list[Path]:
    with tempfile.TemporaryDirectory(prefix="manual-docx-render-") as temporary_dir:
        pdf_path = _convert_pdf(docx_path, Path(temporary_dir))
        return _render_and_verify_pdf(pdf_path, output_dir)


def _convert_pdf(docx_path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("未找到 soffice，无法执行 DOCX 版面验收")
    output_dir.mkdir(parents=True, exist_ok=True)
    environment = os.environ.copy()
    project_fonts = _project_font_dir()
    cjk_font = project_fonts / "NotoSansCJKsc-Regular.otf"
    if not cjk_font.exists() or cjk_font.stat().st_size < 10_000_000:
        raise RuntimeError("缺少完整的 .fonts/NotoSansCJKsc-Regular.otf 中文字体")
    if project_fonts.exists():
        cache_dir = Path(tempfile.mkdtemp(prefix="manual-font-cache-"))
        font_config = output_dir / "fontconfig.xml"
        font_config.write_text(
            "<?xml version=\"1.0\"?>\n"
            "<!DOCTYPE fontconfig SYSTEM \"fonts.dtd\">\n"
            "<fontconfig>\n"
            f"  <dir>{project_fonts.resolve()}</dir>\n"
            f"  <cachedir>{cache_dir}</cachedir>\n"
            "</fontconfig>\n",
            encoding="utf-8",
        )
        environment["FONTCONFIG_FILE"] = str(font_config.resolve())
        environment["FONTCONFIG_PATH"] = str(output_dir.resolve())
        environment["XDG_CACHE_HOME"] = str(cache_dir)
    profile = Path(tempfile.mkdtemp(prefix="manual-lo-profile-"))
    try:
        result = subprocess.run(
            [soffice, "--headless", f"-env:UserInstallation=file://{profile}", "--convert-to", "pdf", "--outdir", str(output_dir), str(docx_path)],
            capture_output=True,
            text=True,
            timeout=120,
            check=False,
            env=environment,
        )
    finally:
        shutil.rmtree(profile, ignore_errors=True)
        if project_fonts.exists():
            shutil.rmtree(cache_dir, ignore_errors=True)
    pdf = output_dir / f"{docx_path.stem}.pdf"
    if result.returncode or not pdf.exists():
        detail = (result.stderr or result.stdout).strip()
        raise RuntimeError(f"LibreOffice PDF 转换失败：{detail or '未生成文件'}")
    return pdf


def _render_and_verify_pdf(pdf_path: Path, output_dir: Path) -> list[Path]:
    pdftoppm = shutil.which("pdftoppm")
    pdfinfo = shutil.which("pdfinfo")
    if not pdftoppm or not pdfinfo:
        raise RuntimeError("未找到 pdftoppm/pdfinfo，无法执行 PDF 逐页验收")
    output_dir.mkdir(parents=True, exist_ok=True)
    for stale_page in output_dir.glob("page-*.png"):
        stale_page.unlink()
    prefix = output_dir / "page"
    result = subprocess.run(
        [pdftoppm, "-png", "-r", "120", str(pdf_path), str(prefix)],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    pages = sorted(output_dir.glob("page-*.png"))
    if result.returncode or not pages:
        raise RuntimeError(f"PDF 逐页渲染失败：{result.stderr.strip()}")
    info = subprocess.run(
        [pdfinfo, str(pdf_path)],
        capture_output=True,
        text=True,
        timeout=30,
        check=False,
    )
    page_line = next(
        (line for line in info.stdout.splitlines() if line.startswith("Pages:")),
        "",
    )
    expected_pages = int(page_line.split(":", 1)[1]) if page_line else 0
    if info.returncode or expected_pages != len(pages):
        raise RuntimeError(
            f"PDF 页数与渲染结果不一致：PDF {expected_pages} 页，PNG {len(pages)} 页"
        )
    for page in pages:
        with Image.open(page) as image:
            grayscale = image.convert("L")
            extrema = grayscale.getextrema()
            if extrema is None or extrema[0] == extrema[1]:
                raise RuntimeError(f"PDF 页面为空或无法识别：{page.name}")
    return pages


def safe_name(value: str) -> str:
    return "".join(char for char in value if char not in '/\\:*?"<>|').strip() or "项目"
